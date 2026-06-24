"""Post-processing: add bookmarks to References entries and hyperlinks to
in-text citations in the body of Research_Proposal_PPRS_Full.docx.

Run this from the project folder where Research_Proposal_PPRS_Full.docx lives:

    cd "C:\\Users\\Laksh\\Desktop\\stock prediction\\binance stock prediction"
    python link_citations.py

It will write Research_Proposal_PPRS_Linked.docx alongside the original,
with every in-text citation in the body converted into a clickable
hyperlink that jumps to the corresponding entry in the References chapter.

Requires python-docx, which you already have installed in your .venv:
    .venv\\Scripts\\activate
    python link_citations.py
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ============================================================================
# Citation map — (bookmark name, author-prefix-for-references-detection,
#                 [list of in-text strings to convert to hyperlinks])
# ============================================================================
CITATIONS = [
    ("ref_angelopoulos_2023", "Angelopoulos",
        ["Angelopoulos and Bates, 2023", "Angelopoulos and Bates (2023)"]),
    ("ref_araci_2019", "Araci",
        ["Araci, 2019", "Araci (2019)"]),
    ("ref_atsalakis_2009", "Atsalakis",
        ["Atsalakis and Valavanis, 2009", "Atsalakis and Valavanis (2009)",
         "Atsalakis & Valavanis (2009)"]),
    ("ref_baltrusaitis_2018", "Baltru",
        ["Baltrušaitis, Ahuja and Morency, 2018",
         "Baltrušaitis, Ahuja and Morency (2018)"]),
    ("ref_bollen_2011", "Bollen",
        ["Bollen, Mao and Zeng, 2011", "Bollen, Mao and Zeng (2011)"]),
    ("ref_bollerslev_1986", "Bollerslev",
        ["Bollerslev, 1986", "Bollerslev (1986)"]),
    ("ref_box_jenkins_1970", "Box",
        ["Box and Jenkins, 1970", "Box and Jenkins (1970)"]),
    ("ref_chen_2016", "Chen",
        ["Chen and Guestrin, 2016", "Chen and Guestrin (2016)"]),
    ("ref_devlin_2019", "Devlin",
        ["Devlin et al., 2019", "Devlin et al. (2019)"]),
    ("ref_diebold_mariano_1995", "Diebold",
        ["Diebold and Mariano, 1995", "Diebold and Mariano (1995)"]),
    ("ref_dosovitskiy_2021", "Dosovitskiy",
        ["Dosovitskiy et al., 2021", "Dosovitskiy et al. (2021)"]),
    ("ref_elkulako_2022", "ElKulako",
        ["ElKulako, 2022", "ElKulako (2022)"]),
    ("ref_harvey_1997", "Harvey",
        ["Harvey, Leybourne and Newbold, 1997",
         "Harvey, Leybourne and Newbold (1997)"]),
    ("ref_he_2016", "He",
        ["He et al., 2016", "He et al. (2016)"]),
    ("ref_hochreiter_1997", "Hochreiter",
        ["Hochreiter and Schmidhuber, 1997", "Hochreiter and Schmidhuber (1997)"]),
    ("ref_johnson_2019", "Johnson",
        ["Johnson et al., 2019", "Johnson et al. (2019)"]),
    ("ref_khedr_2021", "Khedr",
        ["Khedr et al., 2021", "Khedr et al. (2021)"]),
    ("ref_kingma_2015", "Kingma",
        ["Kingma and Ba, 2015", "Kingma and Ba (2015)"]),
    ("ref_kirkpatrick_2017", "Kirkpatrick",
        ["Kirkpatrick et al., 2017", "Kirkpatrick et al. (2017)"]),
    ("ref_livieris_2020", "Livieris",
        ["Livieris, Pintelas and Pintelas, 2020",
         "Livieris, Pintelas and Pintelas (2020)",
         "Livieris et al., 2020", "Livieris et al. (2020)"]),
    ("ref_lo_2004", "Lo",
        ["Lo, 2004", "Lo (2004)"]),
    ("ref_mclean_2016", "McLean",
        ["McLean and Pontiff, 2016", "McLean and Pontiff (2016)"]),
    ("ref_nakamoto_2008", "Nakamoto",
        ["Nakamoto, 2008", "Nakamoto (2008)"]),
    ("ref_nie_2023", "Nie",
        ["Nie et al., 2023", "Nie et al. (2023)"]),
    ("ref_owens_2018", "Owens",
        ["Owens and Efros, 2018", "Owens and Efros (2018)"]),
    ("ref_pelka_2018", "Pelka",
        ["Pelka et al., 2018", "Pelka et al. (2018)"]),
    ("ref_saunders_2019", "Saunders",
        ["Saunders, Lewis and Thornhill, 2019",
         "Saunders, Lewis and Thornhill (2019)"]),
    ("ref_sebastiao_2021", "Sebasti",
        ["Sebastião and Godinho, 2021", "Sebastião and Godinho (2021)"]),
    ("ref_selvaraju_2017", "Selvaraju",
        ["Selvaraju et al., 2017", "Selvaraju et al. (2017)"]),
    ("ref_sezer_2018", "Sezer",
        ["Sezer and Ozbayoglu, 2018", "Sezer and Ozbayoglu (2018)",
         "Sezer & Ozbayoglu (2018)"]),
    ("ref_vaswani_2017", "Vaswani",
        ["Vaswani et al., 2017", "Vaswani et al. (2017)"]),
    ("ref_vovk_2005", "Vovk",
        ["Vovk, Gammerman and Shafer, 2005",
         "Vovk, Gammerman and Shafer (2005)"]),
    ("ref_wang_oates_2015", "Wang",
        ["Wang and Oates, 2015", "Wang and Oates (2015)",
         "Wang & Oates (2015)"]),
    ("ref_wu_2021", "Wu",
        ["Wu et al., 2021", "Wu et al. (2021)"]),
    ("ref_zhou_2021", "Zhou",
        ["Zhou et al., 2021", "Zhou et al. (2021)"]),
]

TEXT_TO_BOOKMARK: dict[str, str] = {}
for bm, _author, patterns in CITATIONS:
    for p_text in patterns:
        TEXT_TO_BOOKMARK[p_text] = bm

ALL_PATTERNS = sorted(TEXT_TO_BOOKMARK.keys(), key=len, reverse=True)
AUTHOR_TO_BOOKMARK = {author_prefix: bm for bm, author_prefix, _ in CITATIONS}


def _bookmark_id_counter():
    n = 1000
    while True:
        yield n
        n += 1


_id_gen = _bookmark_id_counter()


def add_bookmark_around_paragraph(paragraph, name: str):
    bm_id = next(_id_gen)
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bm_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bm_id))
    p_el = paragraph._p
    p_el.insert(0, start)
    p_el.append(end)


def _make_text_run(text: str, *, bold: bool = False, size_pt: int = 12):
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    for tag in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(tag), "Times New Roman")
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(size_pt * 2))
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(size_pt * 2))
    rPr.append(szCs)
    if bold:
        b = OxmlElement("w:b")
        rPr.append(b)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    return r


def _make_hyperlink(text: str, anchor: str, *, size_pt: int = 12):
    hyper = OxmlElement("w:hyperlink")
    hyper.set(qn("w:anchor"), anchor)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    for tag in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(tag), "Times New Roman")
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(size_pt * 2))
    rPr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    hyper.append(r)
    return hyper


def find_citation_spans(text: str):
    matches = []
    consumed = [False] * len(text)
    for pat in ALL_PATTERNS:
        start = 0
        while True:
            idx = text.find(pat, start)
            if idx == -1:
                break
            end = idx + len(pat)
            if any(consumed[idx:end]):
                start = idx + 1
                continue
            matches.append((idx, end, pat))
            for i in range(idx, end):
                consumed[i] = True
            start = end
    matches.sort(key=lambda m: m[0])
    return matches


def rewrite_paragraph_with_links(paragraph) -> bool:
    full_text = paragraph.text
    spans = find_citation_spans(full_text)
    if not spans:
        return False
    size_pt = 12
    bold = False
    if paragraph.runs:
        first_run = paragraph.runs[0]
        if first_run.font.size is not None:
            size_pt = int(first_run.font.size.pt)
        if first_run.font.bold:
            bold = True
    p_el = paragraph._p
    remove = []
    for child in list(p_el):
        tag = child.tag
        if not (tag.endswith("}pPr") or tag.endswith("}bookmarkStart")
                or tag.endswith("}bookmarkEnd")):
            remove.append(child)
    for c in remove:
        p_el.remove(c)
    pos = 0
    for start, end, pat in spans:
        if start > pos:
            p_el.append(_make_text_run(full_text[pos:start], bold=bold, size_pt=size_pt))
        bookmark = TEXT_TO_BOOKMARK[pat]
        p_el.append(_make_hyperlink(pat, bookmark, size_pt=size_pt))
        pos = end
    if pos < len(full_text):
        p_el.append(_make_text_run(full_text[pos:], bold=bold, size_pt=size_pt))
    return True


def process_document(in_path: Path, out_path: Path) -> None:
    doc = Document(in_path)

    in_references = False
    bookmarks_added = 0
    for para in doc.paragraphs:
        txt = (para.text or "").strip()
        if not txt:
            continue
        if txt.upper().startswith("CHAPTER 07") or txt.upper() == "REFERENCES":
            in_references = True
            continue
        if txt.upper().startswith("CHAPTER 08") or txt.upper() == "APPENDICES":
            in_references = False
            continue
        if not in_references:
            continue
        for prefix, bookmark in AUTHOR_TO_BOOKMARK.items():
            if txt.startswith(prefix):
                add_bookmark_around_paragraph(para, bookmark)
                bookmarks_added += 1
                break
    print(f"Bookmarked {bookmarks_added} reference entries.")

    in_references = False
    rewritten = 0
    for para in doc.paragraphs:
        txt = (para.text or "").strip()
        if txt.upper().startswith("CHAPTER 07") or txt.upper() == "REFERENCES":
            in_references = True
            continue
        if txt.upper().startswith("CHAPTER 08") or txt.upper() == "APPENDICES":
            in_references = False
            continue
        if in_references:
            continue
        if rewrite_paragraph_with_links(para):
            rewritten += 1
    print(f"Rewrote {rewritten} body paragraphs with hyperlinks.")

    doc.save(out_path)
    print(f"Saved {out_path}")


if __name__ == "__main__":
    here = Path(__file__).resolve().parent
    in_path = here / "Research_Proposal_PPRS_Full.docx"
    out_path = here / "Research_Proposal_PPRS_Linked.docx"
    if not in_path.exists():
        raise SystemExit(f"Not found: {in_path}")
    process_document(in_path, out_path)
